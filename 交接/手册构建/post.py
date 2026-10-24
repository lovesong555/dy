import docx,sys
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt,Cm
d=docx.Document(sys.argv[1])
def el(tag,**a):
    e=OxmlElement(tag)
    for k,v in a.items(): e.set(qn('w:'+k),v)
    return e
for t in d.tables:
    tblPr=t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b=el('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b.append(el('w:'+side,val='single',sz='4',space='0',color='C9D3E6'))
    tblPr.append(b)
    w=tblPr.find(qn('w:tblW'))
    if w is None: w=el('w:tblW'); tblPr.append(w)
    w.set(qn('w:w'),'5000'); w.set(qn('w:type'),'pct')
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            tcPr=c._tc.get_or_add_tcPr()
            if ri==0: tcPr.append(el('w:shd',val='clear',color='auto',fill='E6EEFB'))
            mar=el('w:tcMar')
            for s in ['top','bottom']: mar.append(el('w:'+s,w='50',type='dxa'))
            for s in ['left','right']: mar.append(el('w:'+s,w='80',type='dxa'))
            tcPr.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
                for r in p.runs:
                    r.font.size=Pt(9)
                    if ri==0: r.font.bold=True
for p in d.paragraphs:
    if p.style.name=='Block Text':
        pPr=p._p.get_or_add_pPr()
        bd=el('w:pBdr'); bd.append(el('w:left',val='single',sz='18',space='8',color='3D8BFF')); pPr.append(bd)
        pPr.append(el('w:shd',val='clear',color='auto',fill='F3F7FF'))
        for r in p.runs:
            if r.text.startswith('【'): r.font.bold=True
# page margins
for s in d.sections:
    s.left_margin=s.right_margin=Cm(2); s.top_margin=s.bottom_margin=Cm(2)
d.save(sys.argv[2])
